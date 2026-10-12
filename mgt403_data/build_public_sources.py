import csv, io, json, os, re, sys, zipfile
from pathlib import Path
import requests
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
RAW = ROOT / "raw"
OUT.mkdir(parents=True, exist_ok=True)
RAW.mkdir(parents=True, exist_ok=True)

YEARS = list(range(2015, 2023))
HEADERS = {"User-Agent": "MGT403 Yale SOM academic project varun.matta@yale.edu"}

CPI = {
    2015: 237.017,
    2016: 240.007,
    2017: 245.120,
    2018: 251.107,
    2019: 255.657,
    2020: 258.811,
    2021: 270.970,
    2022: 292.655,
}

ACS_CODES = [
    "B19301_001E", "B01001_001E", "B23008_002E",
    "B23008_004E", "B23008_010E", "B23008_013E",
    "B25064_001E",
    "B15002_032E", "B15002_033E", "B15002_034E", "B15002_035E", "B15002_019E",
    "B23022_029E", "B23022_026E"
]

def download(url, path, fallbacks=()):
    if path.exists() and path.stat().st_size > 1000:
        return path
    last = None
    for u in (url, *fallbacks):
        try:
            print("Downloading", u, flush=True)
            with requests.get(u, headers=HEADERS, timeout=120, stream=True) as r:
                r.raise_for_status()
                with open(path, "wb") as f:
                    for chunk in r.iter_content(1024 * 1024):
                        if chunk:
                            f.write(chunk)
            if path.stat().st_size > 1000:
                return path
        except Exception as e:
            print("Download failed:", e, flush=True)
            last = e
            if path.exists():
                path.unlink()
    raise RuntimeError(f"Could not download {url}: {last}")

def clean_num(x):
    if x is None:
        return None
    s = str(x).strip().replace(",", "").replace("$", "")
    if s in {"", "*", "#", "**", "NA", "nan", "None"}:
        return None
    try:
        return float(s)
    except Exception:
        return None

def build_ndcp():
    path = download(
        "https://www.dol.gov/sites/dolgov/files/WB/NDCP2022.xlsx",
        RAW / "NDCP2022.xlsx",
        fallbacks=[
            "https://raw.githubusercontent.com/emilylmitchum/dsan5650-final-project/main/assets/NDCP2022.xlsx"
        ]
    )
    print("Reading NDCP workbook", path, path.stat().st_size, flush=True)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [str(x).strip() if x is not None else "" for x in next(it)]
    pos = {name: i for i, name in enumerate(header)}
    aliases = {
        "state_name": ["STATE_NAME", "State_Name"],
        "state_abbr": ["STATE_ABBREVIATION", "State_Abbreviation"],
        "county_name": ["COUNTY_NAME", "County_Name"],
        "county_fips": ["COUNTY_FIPS_CODE", "County_FIPS_Code"],
        "year": ["STUDYYEAR", "StudyYear"],
        "price": ["MCPRESCHOOL", "MCPreschool"],
        "flag": ["MCPreschool_flag", "MCPRESCHOOL_FLAG"],
    }
    def pick(names, required=True):
        for n in names:
            if n in pos: return pos[n]
        if required: raise KeyError(f"Missing any of {names}; sample headers={header[:30]}")
        return None
    p = {k: pick(v, required=(k != "flag")) for k,v in aliases.items()}
    rows = []
    for row in it:
        y = row[p["year"]]
        try: y = int(float(y))
        except Exception: continue
        if y not in YEARS: continue
        fips_raw = row[p["county_fips"]]
        try: fips = str(int(float(fips_raw))).zfill(5)
        except Exception: continue
        price = clean_num(row[p["price"]])
        flag = row[p["flag"]] if p["flag"] is not None else None
        rows.append({
            "year": y,
            "county_fips": fips,
            "state_abbr": row[p["state_abbr"]],
            "state_name": row[p["state_name"]],
            "county_name": row[p["county_name"]],
            "preschool_price_weekly_nominal": price,
            "preschool_price_flag": flag,
        })
    df = pd.DataFrame(rows).drop_duplicates(["year","county_fips"])
    df["cpi_u"] = df["year"].map(CPI)
    df["preschool_price_weekly_2022"] = df["preschool_price_weekly_nominal"] * (CPI[2022] / df["cpi_u"])
    out = OUT / "ndcp_preschool_2015_2022.csv"
    df.to_csv(out, index=False)
    print("NDCP rows", len(df), "nonmissing price", df.preschool_price_weekly_nominal.notna().sum(), flush=True)
    return df

def locate_oews_table(extract_dir):
    files = list(extract_dir.rglob("*"))
    candidates = [p for p in files if p.is_file() and p.suffix.lower() in {".xlsx",".xls"}]
    # prefer metropolitan file, but accept any area file containing MSA
    ranked = sorted(candidates, key=lambda p: (
        0 if re.search(r"MSA_M\d{4}_dl", p.name, re.I) else
        1 if "msa" in p.name.lower() else
        2 if "area" in p.name.lower() else 9,
        len(p.name)
    ))
    for p in ranked:
        if p.name.startswith("~$") or "description" in p.name.lower(): continue
        try:
            test = pd.read_excel(p, nrows=5)
            cols = {str(c).strip().upper() for c in test.columns}
            if "OCC_CODE" in cols and "AREA" in cols and ("A_MEAN" in cols or "ANNUAL_MEAN" in cols):
                return p
        except Exception:
            pass
    raise RuntimeError(f"No OEWS area workbook found in {extract_dir}; files={[p.name for p in candidates][:30]}")

def build_oews():
    all_rows = []
    for year in YEARS:
        yy = str(year)[-2:]
        zpath = download(f"https://www.bls.gov/oes/special-requests/oesm{yy}ma.zip", RAW / f"oesm{yy}ma.zip")
        ex = RAW / f"oesm{yy}ma"
        ex.mkdir(exist_ok=True)
        with zipfile.ZipFile(zpath) as z:
            z.extractall(ex)
        table = locate_oews_table(ex)
        print("OEWS", year, table, flush=True)
        df = pd.read_excel(table, dtype=str)
        df.columns = [str(c).strip().upper() for c in df.columns]
        if "OCC_CODE" not in df.columns:
            raise RuntimeError(f"OCC_CODE missing in {table}")
        sub = df[df["OCC_CODE"].astype(str).str.strip().eq("39-9011")].copy()
        wage_col = "A_MEAN" if "A_MEAN" in sub.columns else "ANNUAL_MEAN"
        sub["a_mean"] = sub[wage_col].map(clean_num)
        sub["area5"] = sub["AREA"].astype(str).str.extract(r"(\d+)")[0].str.zfill(5)
        sub["oews_area"] = "00" + sub["area5"]
        outcols = ["oews_area","a_mean"]
        if "AREA_TITLE" in sub.columns: outcols.append("AREA_TITLE")
        elif "AREA_NAME" in sub.columns: outcols.append("AREA_NAME")
        tmp = sub[outcols].copy()
        tmp["year"] = year
        tmp = tmp.rename(columns={"a_mean":"childcare_worker_wage_nominal","AREA_TITLE":"oews_area_title_bls","AREA_NAME":"oews_area_title_bls"})
        all_rows.append(tmp[["year","oews_area","oews_area_title_bls","childcare_worker_wage_nominal"]] if "oews_area_title_bls" in tmp.columns else tmp[["year","oews_area","childcare_worker_wage_nominal"]])
    out = pd.concat(all_rows, ignore_index=True).drop_duplicates(["year","oews_area"])
    out["cpi_u"] = out["year"].map(CPI)
    out["childcare_worker_wage_2022"] = out["childcare_worker_wage_nominal"] * (CPI[2022] / out["cpi_u"])
    out.to_csv(OUT / "oews_childcare_wages_2015_2022.csv", index=False)
    print("OEWS rows", len(out), "nonmissing", out.childcare_worker_wage_nominal.notna().sum(), flush=True)
    return out

def n(x):
    try:
        v=int(x)
        return v if v >= 0 else None
    except Exception:
        return None

def ratio(a,b):
    return a/b if a is not None and b not in (None,0) else None

def build_acs():
    rows=[]
    for year in YEARS:
        params = {
            "get": "NAME," + ",".join(ACS_CODES),
            "for":"county:*",
            "in":"state:*",
        }
        url=f"https://api.census.gov/data/{year}/acs/acs5"
        print("ACS",year,flush=True)
        r=requests.get(url,params=params,headers=HEADERS,timeout=120)
        r.raise_for_status()
        data=r.json()
        hdr=data[0]
        for rec in data[1:]:
            d=dict(zip(hdr,rec))
            vals={c:n(d.get(c)) for c in ACS_CODES}
            kids_lf = sum([vals["B23008_004E"],vals["B23008_010E"],vals["B23008_013E"]]) if all(vals[c] is not None for c in ["B23008_004E","B23008_010E","B23008_013E"]) else None
            college = sum([vals["B15002_032E"],vals["B15002_033E"],vals["B15002_034E"],vals["B15002_035E"]]) if all(vals[c] is not None for c in ["B15002_032E","B15002_033E","B15002_034E","B15002_035E"]) else None
            fips=str(d["state"]).zfill(2)+str(d["county"]).zfill(3)
            rows.append({
                "year":year,
                "county_fips":fips,
                "county_name_acs":d["NAME"],
                "per_capita_income_nominal":vals["B19301_001E"],
                "pop_total":vals["B01001_001E"],
                "kids_under6":vals["B23008_002E"],
                "kids_under6_allpar_lf":kids_lf,
                "median_gross_rent_nominal":vals["B25064_001E"],
                "female_college_share":ratio(college, vals["B15002_019E"]),
                "female_ftyr_share":ratio(vals["B23022_029E"], vals["B23022_026E"]),
                "under6_share":ratio(vals["B23008_002E"], vals["B01001_001E"]),
                "under6_careneed_share":ratio(kids_lf, vals["B01001_001E"]),
            })
    df=pd.DataFrame(rows)
    df["cpi_u"]=df["year"].map(CPI)
    for col in ["per_capita_income_nominal","median_gross_rent_nominal"]:
        df[col.replace("_nominal","_2022")] = df[col] * (CPI[2022] / df["cpi_u"])
    df.to_csv(OUT/"acs_county_2015_2022.csv",index=False)
    print("ACS rows",len(df),flush=True)
    return df

def main():
    ndcp=build_ndcp()
    oews=build_oews()
    acs=build_acs()
    pd.DataFrame([{"year":y,"cpi_u":v,"to_2022_factor":CPI[2022]/v} for y,v in CPI.items()]).to_csv(OUT/"cpi_u_annual_2015_2022.csv",index=False)
    manifest={
        "ndcp_rows":len(ndcp),"oews_rows":len(oews),"acs_rows":len(acs),
        "years":YEARS,
        "sources":{
            "ndcp":"https://www.dol.gov/sites/dolgov/files/WB/NDCP2022.xlsx",
            "oews":"https://www.bls.gov/oes/special-requests/oesmYYma.zip",
            "acs":"https://api.census.gov/data/YYYY/acs/acs5",
            "cpi":"BLS CPI-U CUUR0000SA0 annual averages"
        }
    }
    (OUT/"manifest.json").write_text(json.dumps(manifest,indent=2))
    print(json.dumps(manifest,indent=2))

if __name__=="__main__":
    main()
